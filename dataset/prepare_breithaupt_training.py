"""Prepare Breithaupt clause-level training data."""

import json
from pathlib import Path

from openpyxl import load_workbook


TARGETS_PATH = Path(
    "data/processed/breithaupt_targets.json"
)

WORKBOOK_PATH = Path(
    r"C:\Datasets\jr2py-osfstorage-archive"
    r"\Stories and retellings"
    r"\All original stories and retellings.xlsx"
)

SHEET_NAME = "Chat and Man"

OUTPUT_PATH = Path(
    "data/processed/breithaupt_training.jsonl"
)


def load_original_stories():
    """Load original stories from the Breithaupt workbook."""

    workbook = load_workbook(
        WORKBOOK_PATH,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook[
        SHEET_NAME
    ]

    stories = {}

    for row in worksheet.iter_rows(
        min_row=3,
        values_only=True,
    ):
        # Column B = Original Story #
        story_id = row[1]

        # Column C = ORIGINAL STORY (human)
        original_story = row[2]

        if (
            story_id is None
            or original_story is None
        ):
            continue

        stories[int(story_id)] = str(
            original_story
        )

    workbook.close()

    return stories


def main():
    with open(
        TARGETS_PATH,
        "r",
        encoding="utf-8",
    ) as f:
        targets = json.load(f)

    original_stories = load_original_stories()

    total_examples = 0

    OUTPUT_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with open(
        OUTPUT_PATH,
        "w",
        encoding="utf-8",
    ) as f:
        for story in targets["stories"]:
            story_id = int(
                story["story_id"]
            )

            if story_id not in original_stories:
                raise ValueError(
                    f"Original story not found "
                    f"for story_id={story_id}"
                )

            paragraph = original_stories[
                story_id
            ]

            for clause in story["clauses"]:
                example = {
                    "narrative_id": story_id,
                    "paragraph": paragraph,
                    "target_clause": clause[
                        "clause"
                    ],
                    "memorability": clause[
                        "memorability"
                    ],
                }

                f.write(
                    json.dumps(
                        example,
                        ensure_ascii=False,
                    )
                    + "\n"
                )

                total_examples += 1

    print(
        f"Stories: "
        f"{len(targets['stories'])}"
    )

    print(
        f"Training examples: "
        f"{total_examples}"
    )

    print(
        f"Saved to: "
        f"{OUTPUT_PATH}"
    )


if __name__ == "__main__":
    main()